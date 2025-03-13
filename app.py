import streamlit as st
import requests
import sqlite3
from requests_oauthlib import OAuth2Session
from PIL import Image
from io import BytesIO
from docx import Document
import os

# Database Setup
DB_FILE = "users.db"

def init_db():
    if not os.path.exists(DB_FILE):
        conn = sqlite3.connect(DB_FILE)
        cursor = conn.cursor()
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                email TEXT UNIQUE NOT NULL,
                role TEXT NOT NULL
            );
        """)
        conn.commit()
        conn.close()

init_db()  # Ensure DB is initialized

def save_user(name, email, role):
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("INSERT OR IGNORE INTO users (name, email, role) VALUES (?, ?, ?)", (name, email, role))
    conn.commit()
    conn.close()

def get_role(email):
    """Determine if the email belongs to a Student or Faculty"""
    if email.startswith("220") and email.endswith("@kiit.ac.in"):
        return "Student"
    elif email.endswith("@kiit.ac.in"):
        return "Faculty"
    return "Unknown"

# OAuth2 Configuration (Use Streamlit Secrets)
CLIENT_ID = st.secrets["oauth"]["client_id"]
CLIENT_SECRET = st.secrets["oauth"]["client_secret"]
REDIRECT_URI = st.secrets["oauth"]["redirect_uri"]

AUTHORIZATION_BASE_URL = "https://accounts.google.com/o/oauth2/auth"
TOKEN_URL = "https://oauth2.googleapis.com/token"
USER_INFO_URL = "https://www.googleapis.com/oauth2/v1/userinfo"

# Certificate Template
TEMPLATE_PATH = "cert.docx"

def generate_certificate(name, role):
    """Generate a certificate from a Word template and return as a downloadable file."""
    doc = Document(TEMPLATE_PATH)

    # Replace placeholders in the template
    for paragraph in doc.paragraphs:
        paragraph.text = paragraph.text.replace("<<Name>>", name)
        paragraph.text = paragraph.text.replace("<<Role>>", role)

    # Save to a BytesIO stream
    output_stream = BytesIO()
    doc.save(output_stream)
    output_stream.seek(0)

    return output_stream

# Streamlit UI
st.title("🎓 Google Sign-In & Certificate Generator")

if "token" not in st.session_state:
    google = OAuth2Session(CLIENT_ID, redirect_uri=REDIRECT_URI, scope=["openid", "email", "profile"])
    authorization_url, state = google.authorization_url(AUTHORIZATION_BASE_URL, access_type="offline")
    st.session_state["oauth_state"] = state
    st.markdown(f"[🔑 Login with Google]({authorization_url})", unsafe_allow_html=True)

# Handle OAuth Callback
query_params = st.query_params
if "code" in query_params:
    code = query_params["code"]
    google = OAuth2Session(CLIENT_ID, redirect_uri=REDIRECT_URI, state=st.session_state["oauth_state"])

    try:
        token = google.fetch_token(TOKEN_URL, client_secret=CLIENT_SECRET, code=code)
        st.session_state["token"] = token

        # Fetch user info
        user_info = requests.get(USER_INFO_URL, headers={"Authorization": f"Bearer {token['access_token']}"}).json()
        st.session_state["user"] = user_info
    except Exception as e:
        st.error(f"OAuth Error: {e}")

# Step 3: Display User Info, Role & Generate Certificate
if "user" in st.session_state:
    user = st.session_state["user"]
    user_email = user["email"]
    user_name = user["name"]
    user_role = get_role(user_email)

    # Save user in DB
    save_user(user_name, user_email, user_role)

    st.success(f"✅ Logged in as {user_name} ({user_email})")
    st.info(f"**Role: {user_role}**")

    # User profile image
    image_url = user.get("picture", "")
    if image_url:
        response = requests.get(image_url)
        img = Image.open(BytesIO(response.content))
        st.image(img, width=100, caption="Google Profile Picture")

    # Certificate Generation
    st.subheader("🎓 Generate Your Certificate")
    name = st.text_input("Enter Your Name", value=user_name)

    if st.button("Generate Certificate"):
        cert_doc = generate_certificate(name, user_role)
        st.download_button("📄 Download Certificate (Word)", data=cert_doc, file_name="certificate.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
